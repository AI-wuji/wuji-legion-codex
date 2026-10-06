import json
from pathlib import Path
import tempfile
import unittest
from unittest import mock
import zipfile
import xml.etree.ElementTree as ET

from docx import Document

from adapters.p4 import office_document_workflow as workflow
from adapters.p4.office_document_workflow import run, validate_request


ROOT = Path(__file__).resolve().parents[1]


class OfficeDocumentWorkflowTests(unittest.TestCase):
    def test_renderer_finds_standard_install_only_in_child_environment(self):
        with tempfile.TemporaryDirectory(dir=ROOT / ".dev/core-test-workspaces") as temporary:
            program = Path(temporary)
            (program / "soffice.exe").write_bytes(b"unexecuted fixture")
            with mock.patch.object(workflow, "LIBREOFFICE_PROGRAM", program), mock.patch.object(
                    workflow.sys, "platform", "win32"), mock.patch.dict(
                    workflow.os.environ, {"PATH": "existing-path"}, clear=True), mock.patch.object(
                    workflow.shutil, "which", return_value=None) as find:
                environment = workflow._render_environment()
                self.assertEqual(environment["PATH"], str(program) + workflow.os.pathsep + "existing-path")
                self.assertEqual(dict(workflow.os.environ), {"PATH": "existing-path"})
                find.assert_called_once_with("soffice.exe", path="existing-path")

    def test_renderer_preserves_existing_path_and_does_not_invent_missing_install(self):
        with tempfile.TemporaryDirectory(dir=ROOT / ".dev/core-test-workspaces") as temporary:
            with mock.patch.object(workflow, "LIBREOFFICE_PROGRAM", Path(temporary)), mock.patch.object(
                    workflow.sys, "platform", "win32"), mock.patch.dict(
                    workflow.os.environ, {"PATH": "existing-path"}, clear=True):
                for located in (None, "already-selected/soffice.exe"):
                    with self.subTest(located=located), mock.patch.object(workflow.shutil, "which", return_value=located):
                        self.assertEqual(workflow._render_environment(), {"PATH": "existing-path"})
                        self.assertEqual(dict(workflow.os.environ), {"PATH": "existing-path"})

    def test_offline_docx_preserves_black_undecorated_title_and_native_content(self):
        request = validate_request({
            "schema_version": 1,
            "kind": "word",
            "title": "\u65e0\u6781\u519b\u56e2 4.0 & <QA> \"Title\"",
            "paragraphs": [
                "\u963f\u6781\u662f\u552f\u4e00\u6c9f\u901a\u5165\u53e3 & <A> \"quoted\" 'text'.",
                "  \u6b63\u6587 <B> & \u4e2d\u6587\u3002  ",
            ],
        })
        with tempfile.TemporaryDirectory(prefix="offline-docx-", dir=ROOT / ".dev/core-test-workspaces") as temporary:
            path = Path(temporary) / "title-regression.docx"
            with (
                mock.patch.object(workflow.office, "OfficeCliAdapter", side_effect=AssertionError("OfficeCLI is forbidden")) as office_adapter,
                mock.patch.object(workflow, "_run_render", side_effect=AssertionError("Rendering is forbidden")) as render,
                mock.patch.object(workflow.subprocess, "run", side_effect=AssertionError("Subprocesses are forbidden")) as subprocess_run,
                mock.patch.object(workflow.subprocess, "Popen", side_effect=AssertionError("Subprocesses are forbidden")) as subprocess_popen,
            ):
                workflow._write_docx(path, request)
            office_adapter.assert_not_called()
            render.assert_not_called()
            subprocess_run.assert_not_called()
            subprocess_popen.assert_not_called()

            document = Document(path)
            self.assertEqual([paragraph.text for paragraph in document.paragraphs], [request["title"], *request["paragraphs"]])
            self.assertEqual(document.core_properties.title, request["title"])
            self.assertEqual(document.paragraphs[0].style.style_id, "Title")
            self.assertEqual(document.styles["Title"].name, "Title")
            self.assertEqual(str(document.styles["Title"].font.color.rgb), "000000")
            self.assertIs(document.styles["Title"].font.underline, False)
            for paragraph in document.paragraphs[1:]:
                self.assertEqual(paragraph.style.style_id, "Normal")

            namespace = {"word": workflow.WORD_NAMESPACE}
            value_attribute = f"{{{workflow.WORD_NAMESPACE}}}val"
            with zipfile.ZipFile(path) as archive:
                styles = ET.fromstring(archive.read("word/styles.xml"))
                body = ET.fromstring(archive.read("word/document.xml"))
                properties = ET.fromstring(archive.read("docProps/core.xml"))
            title_style = styles.find("word:style[@word:styleId='Title']", namespace)
            title_paragraph = body.find("word:body/word:p", namespace)
            self.assertIsNotNone(title_style)
            self.assertIsNotNone(title_paragraph)
            color = title_style.find("word:rPr/word:color", namespace)
            self.assertIsNotNone(color)
            self.assertEqual(color.attrib, {value_attribute: "000000"})
            self.assertEqual(title_paragraph.find("word:pPr/word:pStyle", namespace).get(value_attribute), "Title")
            for source in (title_style, title_paragraph):
                with self.subTest(element=source.tag):
                    self.assertEqual(source.findall(".//word:pBdr", namespace), [])
                    for color in source.findall(".//word:color", namespace):
                        self.assertEqual(color.attrib, {value_attribute: "000000"})
                    for underline in source.findall(".//word:u", namespace):
                        self.assertEqual(underline.get(value_attribute), "none")
            self.assertEqual(properties.findtext("{http://purl.org/dc/elements/1.1/}title"), request["title"])
            inspected = workflow.inspect_docx(path, request)
            self.assertTrue(inspected["passed"], inspected)
            self.assertEqual(inspected["paragraphs"], [request["title"], *request["paragraphs"]])

    def test_sheet_title_uses_black_bold_font_without_changing_header_contrast(self):
        with tempfile.TemporaryDirectory(dir=ROOT / ".dev/core-test-workspaces") as temporary:
            path = Path(temporary) / "title-contrast.xlsx"
            request = validate_request({"schema_version": 1, "kind": "sheet", "title": "Project Costs", "rows": [
                {"label": "Research", "quantity": 2, "unit_price": "120.50"},
            ]})
            workflow._write_xlsx(path, request)
            with zipfile.ZipFile(path) as archive:
                namespace = {"sheet": workflow.OFFICE_NAMESPACE}
                sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
                styles = ET.fromstring(archive.read("xl/styles.xml"))
            fonts = styles.find("sheet:fonts", namespace)
            formats = styles.find("sheet:cellXfs", namespace)
            for reference, expected_color, expected_fill in (("A1", "FF000000", "0"), ("A3", "FFFFFFFF", "2")):
                with self.subTest(cell=reference):
                    cell = sheet.find(f".//sheet:c[@r='{reference}']", namespace)
                    cell_format = formats[int(cell.get("s"))]
                    font = fonts[int(cell_format.get("fontId"))]
                    self.assertEqual(font.find("sheet:color", namespace).get("rgb"), expected_color)
                    self.assertIsNotNone(font.find("sheet:b", namespace))
                    self.assertEqual(cell_format.get("fillId"), expected_fill)
            self.assertTrue(workflow.inspect_xlsx(path, request)["passed"])

    def test_renderer_uses_unique_installed_bundle_not_retired_version(self):
        with tempfile.TemporaryDirectory(dir=ROOT / ".dev/core-test-workspaces") as temporary:
            root = Path(temporary)
            renderer = root / "current/skills/documents/render_docx.py"
            renderer.parent.mkdir(parents=True)
            renderer.write_text("pass", encoding="utf-8")
            with mock.patch.object(workflow, "DOCX_RENDERER_ROOT", root):
                self.assertEqual(workflow._docx_renderer(), renderer)

    def test_absent_and_ambiguous_renderer_do_not_fall_back_or_claim_libreoffice_missing(self):
        with tempfile.TemporaryDirectory(dir=ROOT / ".dev/core-test-workspaces") as temporary:
            root = Path(temporary)
            with mock.patch.object(workflow, "DOCX_RENDERER_ROOT", root):
                with self.assertRaises(workflow.DocxRenderUnavailable) as absent:
                    workflow._docx_renderer()
                self.assertEqual(absent.exception.reason, "renderer_not_found")
                for version in ("first", "second"):
                    renderer = root / version / "skills/documents/render_docx.py"
                    renderer.parent.mkdir(parents=True)
                    renderer.write_text("pass", encoding="utf-8")
                with self.assertRaises(workflow.DocxRenderUnavailable) as ambiguous:
                    workflow._docx_renderer()
                self.assertEqual(ambiguous.exception.reason, "renderer_selection_required")

    def test_missing_libreoffice_preserves_actual_render_command_and_specific_reason(self):
        with tempfile.TemporaryDirectory(dir=ROOT / ".dev/core-test-workspaces") as temporary:
            workspace = Path(temporary)
            receipt = {}
            completed = workflow.subprocess.CompletedProcess([], 1, b"", b"FileNotFoundError: LibreOffice soffice.exe was not found on PATH")
            with mock.patch.object(workflow, "_docx_renderer", return_value=workspace / "renderer.py"), mock.patch.object(
                    workflow.subprocess, "run", return_value=completed):
                with self.assertRaises(workflow.DocxRenderUnavailable) as unavailable:
                    workflow._run_render(workspace / "existing.docx", workspace, receipt)
            self.assertEqual(unavailable.exception.reason, "libreoffice_not_found")
            self.assertEqual(receipt["commands"][0]["exit_code"], 1)
            self.assertIn("soffice.exe", receipt["commands"][0]["stderr"])

    def test_word_and_sheet_requests_preserve_unicode_and_xml_sensitive_text(self):
        word = validate_request({
            "schema_version": 1,
            "kind": "word",
            "title": "\u65e0\u6781\u519b\u56e2 4.0 & <QA>",
            "paragraphs": ["\u963f\u6781\u662f\u552f\u4e00\u6c9f\u901a\u5165\u53e3\u3002"],
        })
        self.assertEqual(word["title"], "\u65e0\u6781\u519b\u56e2 4.0 & <QA>")
        sheet = validate_request({
            "schema_version": 1,
            "kind": "sheet",
            "title": "\u6210\u672c\u8868",
            "rows": [{"label": "\u7814\u7a76 & <A>", "quantity": 2, "unit_price": "0.00"}],
        })
        self.assertEqual(sheet["rows"][0]["label"], "\u7814\u7a76 & <A>")

    def test_request_schema_rejects_undeclared_permissions_and_unsafe_values(self):
        base = {"schema_version": 1, "kind": "sheet", "title": "Cost", "rows": [{"label": "A", "quantity": 1, "unit_price": "1.00"}]}
        invalid = [
            {**base, "formula": "=NOW()"},
            {**base, "rows": [{"label": "A", "quantity": True, "unit_price": "1.00"}]},
            {**base, "rows": [{"label": "A", "quantity": 1, "unit_price": 1.0}]},
            {**base, "rows": [{"label": "A", "quantity": 1, "unit_price": "1"}]},
            {**base, "rows": [{"label": "A\x00", "quantity": 1, "unit_price": "1.00"}]},
            {**base, "rows": [{"label": "A", "quantity": 1, "unit_price": "1000000.01"}]},
        ]
        for request in invalid:
            with self.subTest(request=request):
                with self.assertRaises(ValueError):
                    validate_request(request)

    def test_actual_sheet_creates_formula_bound_workbook_and_rejects_replay_overwrite(self):
        with tempfile.TemporaryDirectory(prefix="office-document-", dir=ROOT / ".dev/core-test-workspaces") as temporary:
            workspace = Path(temporary)
            request = {"schema_version": 1, "kind": "sheet", "title": "Project Cost Summary", "rows": [
                {"label": "Research design & <A>", "quantity": 2, "unit_price": "120.50"},
                {"label": "Isolation test", "quantity": 3, "unit_price": "80.00"},
            ]}
            result = run(workspace, request)
            self.assertTrue(result["passed"], result)
            self.assertEqual(result["build"]["grand_total"], "481.00")
            self.assertTrue(result["checks"]["officecli_observations_completed"])
            self.assertTrue(result["configuration_protection"]["unchanged"])
            replay = run(workspace, request, Path("replay-receipt.json"))
            self.assertFalse(replay["passed"])
            self.assertIn("existing artifact", replay["failure"]["message"])

    def test_actual_word_preserves_structure_and_records_render_block_without_false_success(self):
        with tempfile.TemporaryDirectory(prefix="office-document-word-", dir=ROOT / ".dev/core-test-workspaces") as temporary:
            workspace = Path(temporary)
            request = {"schema_version": 1, "kind": "word", "title": "Wuji Legion 4.0", "paragraphs": ["Aji is the only entry."]}
            with mock.patch.object(workflow, "_run_render", side_effect=workflow.DocxRenderUnavailable(
                    "libreoffice_not_found", "LibreOffice unavailable in isolated failure fixture")):
                result = run(workspace, request)
            self.assertFalse(result["passed"])
            self.assertTrue(result["checks"]["exact_text"])
            self.assertTrue(result["checks"]["officecli_observations_completed"])
            self.assertEqual(result["failure"]["type"], "render_unavailable")
            self.assertEqual(result["render"]["status"], "blocked_external")
            self.assertEqual(result["render"]["reason"], "libreoffice_not_found")
            self.assertTrue(Path(result["artifacts"]["office_file"]["path"]).is_file())
            self.assertTrue(result["configuration_protection"]["unchanged"])


if __name__ == "__main__":
    unittest.main()
