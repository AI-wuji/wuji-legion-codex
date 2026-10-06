from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from PIL import Image

from adapters.p4 import officecli_adapter as office


ROOT = Path(__file__).resolve().parents[2]
DOCX_RENDERER_ROOT = Path.home() / ".codex/plugins/cache/openai-primary-runtime/documents"
MAX_DOCX_BYTES = 8 * 1024 * 1024
MAX_XLSX_BYTES = 8 * 1024 * 1024
CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
DECIMAL_PRICE = re.compile(r"(?:0|[1-9][0-9]{0,6})\.[0-9]{2}")
OFFICE_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _is_clean_text(value: object, limit: int, field: str) -> str:
    if not isinstance(value, str) or CONTROL_CHARACTERS.search(value) or not value.strip() or len(value) > limit:
        raise ValueError(f"{field} must be non-empty bounded plain text")
    return value


def _exact_keys(value: dict, expected: set[str], field: str) -> None:
    if set(value) != expected:
        raise ValueError(f"{field} contains undeclared or missing fields")


def validate_request(value: object) -> dict[str, object]:
    if not isinstance(value, dict):
        raise ValueError("Office document request must be an object")
    if value.get("schema_version") != 1 or type(value.get("schema_version")) is not int:
        raise ValueError("Office document request schema_version must be integer 1")
    kind = value.get("kind")
    if kind not in {"word", "sheet"}:
        raise ValueError("Office document kind must be word or sheet")
    title = _is_clean_text(value.get("title"), 120, "title")
    if kind == "word":
        _exact_keys(value, {"schema_version", "kind", "title", "paragraphs"}, "word request")
        paragraphs = value.get("paragraphs")
        if not isinstance(paragraphs, list) or not 1 <= len(paragraphs) <= 12:
            raise ValueError("word paragraphs must contain 1..12 entries")
        clean_paragraphs = [_is_clean_text(item, 1000, f"paragraph[{index}]") for index, item in enumerate(paragraphs)]
        return {"schema_version": 1, "kind": "word", "title": title, "paragraphs": clean_paragraphs}
    _exact_keys(value, {"schema_version", "kind", "title", "rows"}, "sheet request")
    rows = value.get("rows")
    if not isinstance(rows, list) or not 1 <= len(rows) <= 20:
        raise ValueError("sheet rows must contain 1..20 entries")
    clean_rows = []
    for index, row in enumerate(rows):
        if not isinstance(row, dict):
            raise ValueError(f"row[{index}] must be an object")
        _exact_keys(row, {"label", "quantity", "unit_price"}, f"row[{index}]")
        label = _is_clean_text(row["label"], 80, f"row[{index}].label")
        quantity = row["quantity"]
        if type(quantity) is not int or not 1 <= quantity <= 100000:
            raise ValueError(f"row[{index}].quantity must be an integer in 1..100000")
        unit_price = row["unit_price"]
        if not isinstance(unit_price, str) or not DECIMAL_PRICE.fullmatch(unit_price):
            raise ValueError(f"row[{index}].unit_price must be a non-negative two-decimal string")
        try:
            price = Decimal(unit_price)
        except InvalidOperation as error:
            raise ValueError(f"row[{index}].unit_price is not decimal") from error
        if price > Decimal("1000000.00"):
            raise ValueError(f"row[{index}].unit_price exceeds 1000000.00")
        clean_rows.append({"label": label, "quantity": quantity, "unit_price": unit_price})
    return {"schema_version": 1, "kind": "sheet", "title": title, "rows": clean_rows}


def _workspace(path: Path) -> Path:
    original = Path(path)
    resolved = original.resolve(strict=True)
    roots = [ROOT / ".dev/legion-task-workspaces", ROOT / ".dev/core-test-workspaces"]
    if original.is_symlink() or not resolved.is_dir() or not any(
            resolved != root.resolve() and resolved.is_relative_to(root.resolve()) for root in roots):
        raise ValueError("Office document execution requires a separate isolated workspace")
    return resolved


def _receipt_path(workspace: Path, output: Path | None) -> Path:
    original = Path(output) if output is not None else Path("office-document-receipt.json")
    candidate = original if original.is_absolute() else workspace / original
    resolved = candidate.resolve()
    if candidate.is_symlink() or resolved.parent != workspace or resolved.suffix != ".json" or resolved.exists():
        raise ValueError("Office receipt must be a new direct-child JSON file")
    return resolved


def _artifact_path(workspace: Path, kind: str) -> Path:
    path = workspace / ("word-output.docx" if kind == "word" else "sheet-output.xlsx")
    if path.exists() or path.is_symlink():
        raise FileExistsError("Office document workflow never overwrites an existing artifact")
    return path


def _qname(namespace: str, name: str) -> str:
    return f"{{{namespace}}}{name}"


def _inline_string(parent: ET.Element, reference: str, value: str, style: int | None = None) -> None:
    attributes = {"r": reference, "t": "inlineStr"}
    if style is not None:
        attributes["s"] = str(style)
    cell = ET.SubElement(parent, _qname(OFFICE_NAMESPACE, "c"), attributes)
    inline = ET.SubElement(cell, _qname(OFFICE_NAMESPACE, "is"))
    text = ET.SubElement(inline, _qname(OFFICE_NAMESPACE, "t"))
    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text.text = value


def _number(parent: ET.Element, reference: str, value: Decimal | int | float, style: int | None = None) -> None:
    attributes = {"r": reference, "t": "n"}
    if style is not None:
        attributes["s"] = str(style)
    cell = ET.SubElement(parent, _qname(OFFICE_NAMESPACE, "c"), attributes)
    value_element = ET.SubElement(cell, _qname(OFFICE_NAMESPACE, "v"))
    value_element.text = str(value)


def _formula(parent: ET.Element, reference: str, expression: str, cached: Decimal | int, style: int | None = None) -> None:
    attributes = {"r": reference}
    if style is not None:
        attributes["s"] = str(style)
    cell = ET.SubElement(parent, _qname(OFFICE_NAMESPACE, "c"), attributes)
    formula = ET.SubElement(cell, _qname(OFFICE_NAMESPACE, "f"))
    formula.text = expression
    value_element = ET.SubElement(cell, _qname(OFFICE_NAMESPACE, "v"))
    value_element.text = str(cached)


def _write_xlsx(path: Path, request: dict[str, object]) -> dict[str, object]:
    rows = request["rows"]
    assert isinstance(rows, list)
    sheet = ET.Element(_qname(OFFICE_NAMESPACE, "worksheet"))
    ET.SubElement(sheet, _qname(OFFICE_NAMESPACE, "sheetPr"))
    ET.SubElement(sheet, _qname(OFFICE_NAMESPACE, "dimension"), {"ref": f"A1:D{len(rows) + 4}"})
    views = ET.SubElement(sheet, _qname(OFFICE_NAMESPACE, "sheetViews"))
    ET.SubElement(views, _qname(OFFICE_NAMESPACE, "sheetView"), {"showGridLines": "0", "workbookViewId": "0"})
    ET.SubElement(sheet, _qname(OFFICE_NAMESPACE, "sheetFormatPr"), {"defaultRowHeight": "18"})
    columns = ET.SubElement(sheet, _qname(OFFICE_NAMESPACE, "cols"))
    for minimum, maximum, width in ((1, 1, "28"), (2, 2, "12"), (3, 3, "15"), (4, 4, "15")):
        ET.SubElement(columns, _qname(OFFICE_NAMESPACE, "col"), {"min": str(minimum), "max": str(maximum), "width": width, "customWidth": "1"})
    sheet_data = ET.SubElement(sheet, _qname(OFFICE_NAMESPACE, "sheetData"))
    row = ET.SubElement(sheet_data, _qname(OFFICE_NAMESPACE, "row"), {"r": "1"})
    _inline_string(row, "A1", str(request["title"]), 3)
    row = ET.SubElement(sheet_data, _qname(OFFICE_NAMESPACE, "row"), {"r": "3"})
    for reference, value in zip(("A3", "B3", "C3", "D3"), ("Item", "Quantity", "Unit Price", "Total")):
        _inline_string(row, reference, value, 1)
    total = Decimal("0.00")
    for index, item in enumerate(rows, start=4):
        assert isinstance(item, dict)
        price = Decimal(str(item["unit_price"]))
        line_total = price * int(item["quantity"])
        total += line_total
        row = ET.SubElement(sheet_data, _qname(OFFICE_NAMESPACE, "row"), {"r": str(index)})
        _inline_string(row, f"A{index}", str(item["label"]))
        _number(row, f"B{index}", int(item["quantity"]))
        _number(row, f"C{index}", f"{price:.2f}", 2)
        _formula(row, f"D{index}", f"B{index}*C{index}", f"{line_total:.2f}", 2)
    total_row = len(rows) + 4
    row = ET.SubElement(sheet_data, _qname(OFFICE_NAMESPACE, "row"), {"r": str(total_row)})
    _inline_string(row, f"A{total_row}", "Grand Total", 1)
    _formula(row, f"D{total_row}", f"SUM(D4:D{total_row - 1})", f"{total:.2f}", 2)
    ET.SubElement(sheet, _qname(OFFICE_NAMESPACE, "pageMargins"), {"left": "0.3", "right": "0.3", "top": "0.5", "bottom": "0.5", "header": "0.3", "footer": "0.3"})

    workbook = ET.Element(_qname(OFFICE_NAMESPACE, "workbook"))
    ET.SubElement(workbook, _qname(OFFICE_NAMESPACE, "fileVersion"), {"appName": "WujiLegion4", "lastEdited": "1", "lowestEdited": "1", "rupBuild": "1"})
    sheets = ET.SubElement(workbook, _qname(OFFICE_NAMESPACE, "sheets"))
    ET.SubElement(sheets, _qname(OFFICE_NAMESPACE, "sheet"), {"name": "Summary", "sheetId": "1", _qname(RELATIONSHIP_NAMESPACE, "id"): "rId1"})
    ET.SubElement(workbook, _qname(OFFICE_NAMESPACE, "calcPr"), {"calcId": "1", "calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"})

    styles = ET.fromstring(f'''<styleSheet xmlns="{OFFICE_NAMESPACE}"><numFmts count="1"><numFmt numFmtId="164" formatCode="0.00"/></numFmts><fonts count="2"><font><sz val="11"/><color theme="1"/><name val="Arial"/></font><font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Arial"/></font></fonts><fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FF1F4E78"/><bgColor indexed="64"/></patternFill></fill></fills><borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="4"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/><xf numFmtId="0" fontId="1" fillId="2" borderId="0" applyFont="1" applyFill="1"/><xf numFmtId="164" fontId="0" fillId="0" borderId="0" applyNumberFormat="1"/><xf numFmtId="0" fontId="1" fillId="0" borderId="0" applyFont="1"/></cellXfs><cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles><dxfs count="0"/><tableStyles count="0"/></styleSheet>''')
    content_types = ET.fromstring('''<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/><Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/><Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/></Types>''')
    root_relationships = ET.fromstring('''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/><Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/></Relationships>''')
    workbook_relationships = ET.fromstring('''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>''')
    core = ET.Element(_qname("http://schemas.openxmlformats.org/package/2006/metadata/core-properties", "coreProperties"))
    ET.SubElement(core, _qname("http://purl.org/dc/elements/1.1/", "title")).text = str(request["title"])
    ET.SubElement(core, _qname("http://purl.org/dc/elements/1.1/", "creator")).text = "Wuji Legion 4.0"
    app = ET.fromstring('''<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"><Application>Wuji Legion 4.0</Application><DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop></Properties>''')
    parts = {"[Content_Types].xml": content_types, "_rels/.rels": root_relationships, "xl/workbook.xml": workbook, "xl/_rels/workbook.xml.rels": workbook_relationships, "xl/worksheets/sheet1.xml": sheet, "xl/styles.xml": styles, "docProps/core.xml": core, "docProps/app.xml": app}
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name, element in parts.items():
            namespace = None
            if name == "[Content_Types].xml":
                namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
            elif name.endswith(".rels"):
                namespace = PACKAGE_RELATIONSHIP_NAMESPACE
            elif name == "docProps/app.xml":
                namespace = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
            elif name.startswith("xl/"):
                namespace = OFFICE_NAMESPACE
            archive.writestr(name, _xml_bytes(element, namespace))
    return {"sheet": "Summary", "data_rows": len(rows), "formula_rows": len(rows) + 1, "grand_total": f"{total:.2f}"}


def _xml_escape(value: str) -> str:
    return value


def _xml_bytes(element: ET.Element, default_namespace: str | None = None) -> bytes:
    if default_namespace is not None:
        ET.register_namespace("", default_namespace)
    ET.register_namespace("r", RELATIONSHIP_NAMESPACE)
    ET.register_namespace("pr", PACKAGE_RELATIONSHIP_NAMESPACE)
    ET.register_namespace("cp", "http://schemas.openxmlformats.org/package/2006/metadata/core-properties")
    ET.register_namespace("dc", "http://purl.org/dc/elements/1.1/")
    return ET.tostring(element, encoding="utf-8", xml_declaration=True)


def inspect_docx(path: Path, request: dict[str, object]) -> dict[str, object]:
    expected = [str(request["title"]), *[str(item) for item in request["paragraphs"]]]
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or sum(item.file_size for item in archive.infolist()) > MAX_DOCX_BYTES or archive.testzip() is not None:
            raise ValueError("DOCX ZIP integrity or size check failed")
        document = ET.fromstring(archive.read("word/document.xml"))
        paragraphs = document.findall(f".//{{{WORD_NAMESPACE}}}p")
        texts = ["".join(text.text or "" for text in paragraph.findall(f".//{{{WORD_NAMESPACE}}}t")) for paragraph in paragraphs]
        styles = [paragraph.find(f"{{{WORD_NAMESPACE}}}pPr/{{{WORD_NAMESPACE}}}pStyle") for paragraph in paragraphs]
        external = []
        if "word/_rels/document.xml.rels" in names:
            relationships = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
            external = [item.get("Target") for item in relationships if item.get("TargetMode", "").casefold() == "external"]
        checks = {
            "exact_text": texts == expected,
            "title_style": bool(styles and styles[0] is not None and styles[0].get(f"{{{WORD_NAMESPACE}}}val") == "Title"),
            "no_external_relationships": not external,
            "no_embedded_or_macro_parts": not any(name.startswith(("word/media/", "word/embeddings/")) or name.lower().endswith("vbaProject.bin") for name in names),
            "package_has_document_xml": "word/document.xml" in names,
        }
    return {"checks": checks, "passed": all(checks.values()), "paragraphs": texts, "external_targets": external}


def _xlsx_cells(path: Path) -> dict[str, dict[str, str]]:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    result = {}
    for cell in sheet.findall(f".//{{{OFFICE_NAMESPACE}}}c"):
        reference = cell.get("r")
        value = cell.find(f"{{{OFFICE_NAMESPACE}}}v")
        formula = cell.find(f"{{{OFFICE_NAMESPACE}}}f")
        text = cell.find(f"{{{OFFICE_NAMESPACE}}}is/{{{OFFICE_NAMESPACE}}}t")
        result[reference] = {"value": (text.text if text is not None else value.text if value is not None else ""), "formula": formula.text if formula is not None else None}
    return result


def inspect_xlsx(path: Path, request: dict[str, object]) -> dict[str, object]:
    cells = _xlsx_cells(path)
    rows = request["rows"]
    assert isinstance(rows, list)
    checks = {
        "title": cells.get("A1", {}).get("value") == request["title"],
        "headers": [cells.get(reference, {}).get("value") for reference in ("A3", "B3", "C3", "D3")] == ["Item", "Quantity", "Unit Price", "Total"],
        "formula_count": sum(1 for cell in cells.values() if cell.get("formula")) == len(rows) + 1,
        "no_external_formulas": all("[" not in str(cell.get("formula")) and "://" not in str(cell.get("formula")) for cell in cells.values()),
    }
    total = Decimal("0.00")
    for index, item in enumerate(rows, start=4):
        assert isinstance(item, dict)
        line_total = Decimal(str(item["unit_price"])) * int(item["quantity"])
        total += line_total
        checks[f"row_{index}"] = (
            cells.get(f"A{index}", {}).get("value") == item["label"]
            and cells.get(f"B{index}", {}).get("value") == str(item["quantity"])
            and cells.get(f"C{index}", {}).get("value") == f"{Decimal(str(item['unit_price'])):.2f}"
            and cells.get(f"D{index}", {}).get("formula") == f"B{index}*C{index}"
        )
    total_row = len(rows) + 4
    checks["grand_total_formula"] = cells.get(f"D{total_row}", {}).get("formula") == f"SUM(D4:D{total_row - 1})"
    return {"checks": checks, "passed": all(checks.values()), "grand_total": f"{total:.2f}", "cells": len(cells)}


class DocxRenderUnavailable(FileNotFoundError):
    def __init__(self, reason: str, message: str):
        super().__init__(message)
        self.reason = reason


def _docx_renderer() -> Path:
    candidates = [path for path in DOCX_RENDERER_ROOT.glob("*/skills/documents/render_docx.py")
                  if path.is_file() and not path.is_symlink()]
    if not candidates:
        raise DocxRenderUnavailable("renderer_not_found", "Approved bundled DOCX renderer is unavailable")
    if len(candidates) != 1:
        raise DocxRenderUnavailable("renderer_selection_required", "Multiple bundled DOCX renderers require explicit runtime selection")
    return candidates[0]


def _run_render(path: Path, workspace: Path, receipt: dict[str, object]) -> dict[str, object]:
    renderer = _docx_renderer()
    render_dir = workspace / "docx-render-qa"
    if render_dir.exists():
        raise FileExistsError(str(render_dir))
    render_dir.mkdir()
    command = [sys.executable, str(renderer), str(path), "--output_dir", str(render_dir)]
    environment = dict(os.environ)
    completed = subprocess.run(command, cwd=workspace, env=environment, capture_output=True, timeout=90,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if len(completed.stdout) + len(completed.stderr) > 256 * 1024:
        raise ValueError("DOCX renderer response exceeded the bounded evidence limit")
    record = {"phase": "docx-render", "argv": command, "exit_code": completed.returncode,
              "stdout": completed.stdout.decode("utf-8", errors="replace"), "stderr": completed.stderr.decode("utf-8", errors="replace")}
    receipt.setdefault("commands", []).append(record)
    if completed.returncode != 0 and "FileNotFoundError:" in record["stderr"] and "LibreOffice" in record["stderr"]:
        raise DocxRenderUnavailable("libreoffice_not_found", record["stderr"].strip())
    if completed.returncode != 0:
        raise RuntimeError("DOCX renderer failed")
    pages = []
    for page in sorted(render_dir.glob("page-*.png")):
        with Image.open(page) as image:
            pages.append({"path": str(page), "bytes": page.stat().st_size, "sha256": digest(page), "width": image.width, "height": image.height})
    if not pages or any(page["bytes"] <= 0 or page["width"] <= 0 or page["height"] <= 0 for page in pages):
        raise RuntimeError("DOCX render produced no valid page images")
    return {"directory": str(render_dir), "pages": pages}


def run(workspace: Path, request: dict[str, object], output: Path | None = None) -> dict[str, object]:
    isolated = _workspace(workspace)
    receipt_path = _receipt_path(isolated, output)
    receipt: dict[str, object] = {"schema_version": 1, "kind": "isolated-office-document-task", "status": "running", "result_state": "running", "passed": False, "receipt_path": str(receipt_path), "commands": [], "checks": {}, "artifacts": {}, "professional_effectiveness": "not_claimed", "P7": False, "shutdown": False, "failure": None, "started_at_utc": utc_now()}
    with receipt_path.open("x", encoding="utf-8") as handle:
        json.dump(receipt, handle, ensure_ascii=False)
    config_before = None
    adapter = None
    artifact_path = None
    try:
        request = validate_request(request)
        receipt["request"] = request
        receipt["request_sha256"] = hashlib.sha256(json.dumps(request, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
        config_before = office.protected_snapshot()
        adapter = office.OfficeCliAdapter(isolated)
        receipt["tool"] = {"version": office.PINNED_VERSION, "binary": str(adapter.binary), "binary_sha256": digest(adapter.binary)}
        artifact_path = _artifact_path(isolated, str(request["kind"]))
        if request["kind"] == "word":
            document = Document()
            document.core_properties.title = str(request["title"])
            document.add_paragraph(str(request["title"]), style="Title")
            for paragraph in request["paragraphs"]:
                document.add_paragraph(str(paragraph), style="Normal")
            document.save(artifact_path)
            inspected = inspect_docx(artifact_path, request)
            receipt["checks"].update(inspected["checks"])
            receipt["artifacts"]["office_file"] = {"path": str(artifact_path), "bytes": artifact_path.stat().st_size, "sha256": digest(artifact_path), "kind": str(request["kind"])}
            adapter.owned.add(artifact_path.name)
            adapter.observe("validate", artifact_path.name)
            adapter.observe("view", artifact_path.name, "outline")
            adapter.observe("get", artifact_path.name, "/body")
            try:
                receipt["render"] = _run_render(artifact_path, isolated, receipt)
                receipt["checks"]["rendered_pages"] = bool(receipt["render"]["pages"])
            except DocxRenderUnavailable as error:
                receipt["render"] = {"status": "blocked_external", "reason": error.reason, "error": str(error)}
                receipt["failure"] = {"type": "render_unavailable", "message": str(error)}
                receipt["checks"]["rendered_pages"] = False
        else:
            receipt["build"] = _write_xlsx(artifact_path, request)
            inspected = inspect_xlsx(artifact_path, request)
            adapter.owned.add(artifact_path.name)
            adapter.observe("validate", artifact_path.name)
            adapter.observe("view", artifact_path.name, "outline")
            adapter.observe("get", artifact_path.name, "/Summary")
        receipt["checks"].update(inspected["checks"])
        receipt["checks"]["artifact_nonempty"] = artifact_path.is_file() and artifact_path.stat().st_size > 0
        receipt["checks"]["officecli_observations_completed"] = all(item.get("exit_code") == 0 for item in adapter.receipts)
        receipt["artifacts"]["office_file"] = {"path": str(artifact_path), "bytes": artifact_path.stat().st_size, "sha256": digest(artifact_path), "kind": str(request["kind"])}
        receipt["tool_observations"] = adapter.receipts
    except Exception as error:
        receipt["failure"] = {"type": type(error).__name__, "message": str(error)}
    finally:
        if adapter is not None:
            receipt["tool_observations"] = adapter.receipts
        if artifact_path is not None and artifact_path.is_file() and "office_file" not in receipt["artifacts"]:
            receipt["artifacts"]["office_file"] = {"path": str(artifact_path), "bytes": artifact_path.stat().st_size, "sha256": digest(artifact_path), "kind": str(request.get("kind", "unknown")) if isinstance(request, dict) else "unknown"}
        try:
            config_after = office.protected_snapshot()
        except Exception:
            config_after = None
        receipt["configuration_protection"] = {"before": config_before, "after": config_after, "unchanged": config_before is not None and config_before == config_after}
        receipt["checks"]["configuration_unchanged"] = receipt["configuration_protection"]["unchanged"]
        receipt["passed"] = receipt["failure"] is None and bool(receipt["checks"]) and all(receipt["checks"].values())
        receipt["status"] = receipt["result_state"] = "completed_bounded_task" if receipt["passed"] else "failed_evidence_retained"
        receipt["finished_at_utc"] = utc_now()
        receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return receipt
